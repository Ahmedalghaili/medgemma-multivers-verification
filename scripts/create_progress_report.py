"""Create a concise professor-facing progress report from recorded project results."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

OUT = "results/medical_claim_verification_progress_report.docx"
DIAGRAM = "results/medical_claim_verification_workflow.png"
NAVY = RGBColor(25, 48, 77)
TEAL = RGBColor(0, 125, 135)
GRAY = RGBColor(90, 102, 115)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.82)
section.right_margin = Inches(.82)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(5)
for name, size in [('Title', 21), ('Heading 1', 13), ('Heading 2', 10.5)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY
    st.paragraph_format.space_before = Pt(11)
    st.paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, '19304D')
        for r in c.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j % 2 == 0:
                shade(cells[i], 'EEF4F7')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def make_diagram():
    im = Image.new('RGB', (1500, 600), '#FFFFFF')
    draw = ImageDraw.Draw(im)
    font = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    bold = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
    title_font = ImageFont.truetype(bold, 29)
    body_font = ImageFont.truetype(font, 25)
    arrow_font = ImageFont.truetype(bold, 34)
    steps = [
        ('1  QUESTION', 'Can face masks reduce virus spread?', '#DCEBF6'),
        ('2  MEDGEMMA WRITES', 'Drafts an answer in everyday language', '#DCEBF6'),
        ('3  SPLIT INTO CLAIMS', 'Example claim: Masks reduce respiratory droplets', '#D9F0EE'),
        ('4  SEARCH RESEARCH', 'Finds relevant biomedical abstracts', '#D9F0EE'),
        ('5  MULTIVERS CHECKS', 'Support / contradict / NEI + evidence sentences', '#E4EBDC'),
    ]
    for i, (heading, detail, fill) in enumerate(steps):
        y = 15 + i * 116
        draw.rounded_rectangle((75, y, 1425, y + 86), radius=20, fill=fill, outline='#B9CBD6', width=3)
        draw.text((113, y + 6), heading, font=title_font, fill='#19304D')
        draw.text((113, y + 47), detail, font=body_font, fill='#354859')
        if i < len(steps) - 1:
            draw.text((730, y + 83), '↓', font=arrow_font, fill='#007D87', anchor='mt')
    im.save(DIAGRAM)

make_diagram()

p = doc.add_paragraph(style='Title')
p.add_run('Medical Claim Verification').font.color.rgb = NAVY
p = doc.add_paragraph()
p.add_run('Project progress report').bold = True
p.add_run('  |  22 September 2026')
p.style = 'Subtitle'

doc.add_heading('Research objective', 1)
doc.add_paragraph('Can we make AI-generated medical answers more trustworthy by checking their statements against research? This project combines MedGemma, which writes the answer, with MultiVerS, which checks the statements.')

doc.add_heading('What is MultiVerS?', 1)
doc.add_paragraph('MultiVerS is an AI evidence checker. Give it one medical statement and one research paper abstract (a short summary of a paper). It reads both and gives a result: SUPPORT, CONTRADICT, or NOT ENOUGH INFORMATION (NEI). It also points to the sentences in the abstract that led to the result. We trained and fine-tuned it using examples from the HealthVer dataset.')

doc.add_heading('Progress at a glance', 1)
table(['Completed', 'Current status'], [
    ('HealthVer data preparation', '5,292 training, 940 validation, and 903 test claim–abstract pairs; 322 unique abstracts.'),
    ('Verifier training', 'Three MultiVerS versions evaluated on the held-out HealthVer test set.'),
    ('Pipeline integration', 'MedGemma answer generation, claim extraction, semantic retrieval, and frozen MultiVerS verification connected.'),
], [2.0, 4.8])

doc.add_heading('System workflow', 1)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(DIAGRAM, width=Inches(6.7))

h = doc.add_heading('Model setup and training', 1)
h.paragraph_format.page_break_before = True
doc.add_paragraph('All three versions use a Longformer science encoder. The input is one claim plus one research abstract. The model has two tasks: choose a claim label (SUPPORT, CONTRADICT, or NEI) and select the abstract sentences used as evidence. Training used 5,292 HealthVer examples; 940 examples were used for validation, and the same 903 examples were kept for final testing.')
table(['Version', 'What changed and why'], [
    ('Original MultiVerS', 'Started with a science-trained Longformer encoder and trained the full MultiVerS model on HealthVer.'),
    ('Fine-tuned MultiVerS', 'Continued training from the best original checkpoint with a smaller learning rate. This lets the model make careful adjustments to what it already learned.'),
    ('Label-aware modified model', 'Started from the fine-tuned checkpoint. The model’s confidence in each possible label was also given to the evidence selector to help it choose sentences.'),
], [1.7, 5.1])
doc.add_paragraph('Training configuration: one GPU; batch size 1 with eight gradient accumulation steps (effective batch size 8); 16-bit precision. Original: learning rate 1×10⁻⁵, up to 20 epochs. Fine-tuning started at 2×10⁻⁶ and resumed from a saved checkpoint to complete up to 5 epochs. Label-aware: 1×10⁻⁵, up to 5 epochs. The best checkpoint was selected using validation sentence-label F1.')

doc.add_heading('Results obtained so far', 1)
doc.add_heading('MultiVerS model results: HealthVer test set', 2)
doc.add_paragraph('These scores test MultiVerS when the claim and research abstract are already provided. Accuracy for the full MedGemma workflow awaits human review.')
table(['Model', 'Accuracy', 'Macro-F1', 'Abstract F1', 'Sentence selection F1', 'Sentence label F1'], [
    ('Original MultiVerS', '74.86%', '74.37%', '74.90%', '84.58%', '73.52%'),
    ('Fine-tuned MultiVerS', '75.30%', '74.89%', '75.43%', '85.41%', '74.82%'),
    ('Label-aware modified model', '74.75%', '74.37%', '74.23%', '82.91%', '73.04%'),
], [1.6, .85, .85, .85, 1.3, 1.2])
doc.add_paragraph('The fine-tuned MultiVerS model scored highest on all five measures, so it was selected for the full workflow. The extra label-aware connection did not improve scores. Accuracy and macro-F1 measure claim labels; the other F1 scores measure evidence-related predictions.')

h = doc.add_heading('What the five measures mean', 2)
h.paragraph_format.page_break_before = True
table(['Measure', 'Simple meaning'], [
    ('Accuracy', 'How often the claim label is correct.'),
    ('Macro-F1', 'How well the model handles all three labels, giving each equal weight.'),
    ('Abstract F1', 'How well it identifies the correct SUPPORT or CONTRADICT result for an abstract.'),
    ('Sentence selection F1', 'How well it finds the correct evidence sentences.'),
    ('Sentence label F1', 'How well it gets both the evidence sentences and their claim label right.'),
], [1.7, 5.1])

doc.add_heading('Interpretation and next steps', 1)
bullet('Evaluate the complete MedGemma workflow using human-reviewed claim labels and evidence.')
bullet('Review errors in answer writing, claim splitting, evidence search, and verification.')
bullet('Compare MedGemma alone with the evidence-verified workflow after human evaluation.')

doc.add_heading('Key takeaway', 1)
p = doc.add_paragraph('A working medical claim verification workflow is in place. Fine-tuned MultiVerS performed best on the HealthVer test set. The full workflow still needs human evaluation before its accuracy can be reported.')
p.runs[0].bold = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Medical Claim Verification  •  Progress Report')
run.font.size = Pt(8)
run.font.color.rgb = GRAY

doc.save(OUT)
print(OUT)
