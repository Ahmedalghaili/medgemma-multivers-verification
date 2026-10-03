from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = "results/medical_claim_verification_project.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(20, 44, 78)
BLUE = RGBColor(35, 112, 180)
TEAL = RGBColor(0, 145, 150)
LIGHT = RGBColor(242, 247, 251)
DARK = RGBColor(38, 48, 60)
GRAY = RGBColor(100, 112, 125)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(34, 130, 85)
ORANGE = RGBColor(210, 117, 35)

def textbox(slide, x, y, w, h, text, size=20, color=DARK, bold=False,
            align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box

def title(slide, text, subtitle=None):
    textbox(slide, .55, .3, 12.2, .55, text, 28, NAVY, True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.0), Inches(1.25), Inches(.08))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    if subtitle: textbox(slide, .55, 1.13, 12, .35, subtitle, 12, GRAY)

def new_slide(text, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    title(s, text, subtitle); return s

def footer(slide, n):
    textbox(slide, .6, 7.12, 11.8, .2, "Medical Claim Verification Project", 9, GRAY)
    textbox(slide, 12.35, 7.12, .35, .2, str(n), 9, GRAY, align=PP_ALIGN.RIGHT)

def bullets(slide, items, x=.8, y=1.65, w=11.7, h=4.9, size=20, color=DARK):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item; p.level=0; p.space_after=Pt(12)
        p.font.name="Aptos"; p.font.size=Pt(size); p.font.color.rgb=color
    return box

def card(slide, x, y, w, h, heading, body, accent=BLUE, body_size=16):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=RGBColor(215,225,235)
    textbox(slide,x+.18,y+.15,w-.36,.35,heading,17,accent,True)
    textbox(slide,x+.18,y+.62,w-.36,h-.78,body,body_size,DARK)

def add_table(slide, rows, cols, data, x, y, w, h, widths=None, font_size=13):
    table=slide.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(h)).table
    if widths:
        for i,val in enumerate(widths): table.columns[i].width=Inches(val)
    for r in range(rows):
        for c in range(cols):
            cell=table.cell(r,c); cell.text=str(data[r][c]); cell.margin_left=Inches(.08); cell.margin_right=Inches(.05)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r==0 else (LIGHT if r%2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name="Aptos"; p.font.size=Pt(font_size); p.font.bold=(r==0); p.font.color.rgb=WHITE if r==0 else DARK
    return table

# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
textbox(s,.8,1.45,11.8,1.2,"Medical Claim Verification",38,WHITE,True,PP_ALIGN.CENTER)
textbox(s,1.3,2.85,10.8,.7,"From HealthVer training to a MedGemma + MultiVerS evidence pipeline",22,RGBColor(200,225,240),False,PP_ALIGN.CENTER)
textbox(s,2,4.25,9.3,.8,"Project progress presentation",20,WHITE,False,PP_ALIGN.CENTER)
textbox(s,2,5.55,9.3,.4,"Dataset • training • evaluation • final system",15,RGBColor(190,205,220),False,PP_ALIGN.CENTER)

# 2
s=new_slide("1. Project goal","Build and evaluate a system that checks medical claims against evidence.")
card(s,.75,1.75,3.7,3.8,"Input","A medical question\n\nExample: What are the health effects of wearing face masks?",BLUE,18)
card(s,4.8,1.75,3.7,3.8,"Verification","Split the answer into claims\n\nFor each claim: SUPPORT, CONTRADICT, or NEI",TEAL,18)
card(s,8.85,1.75,3.7,3.8,"Output","A safer answer with selected evidence sentences and confidence",GREEN,18)
footer(s,2)

# 3
s=new_slide("2. Dataset used: HealthVer","HealthVer is an entailment dataset for checking health-related claims.")
add_table(s,5,2,[["Split / content","Size"],["Training claim-document pairs","5,292"],["Validation pairs","940"],["Test pairs","903"],["Unique abstracts","322"]],.8,1.65,5.5,3.6,[3.8,1.7],16)
card(s,6.8,1.65,5.6,3.6,"Each example contains","• A medical claim\n• One biomedical abstract\n• Evidence sentence numbers\n• A label: SUPPORT, CONTRADICT, or NEI",TEAL,18)
textbox(s,.9,5.65,11.4,.7,"The evidence labels are sentence-level, so the model learns both claim verification and evidence selection.",20,NAVY,True,PP_ALIGN.CENTER)
footer(s,3)

# 4
s=new_slide("3. What I used","Main tools and hardware used for the experiments.")
card(s,.75,1.55,3.8,4.4,"Software","Python 3.12\nPyTorch with CUDA\nTransformers\nPyTorch Lightning\nHugging Face Datasets\nTensorBoard",BLUE,17)
card(s,4.8,1.55,3.8,4.4,"Models","Longformer encoder for MultiVerS\n\nGoogle MedGemma-4B-it for answer generation\n\nS-PubMedBERT-MS-MARCO for semantic retrieval",TEAL,17)
card(s,8.85,1.55,3.8,4.4,"Hardware","NVIDIA RTX 4080 SUPER\n16 GB VRAM\n\nFine-tuned models and checkpoints were saved locally for reproducibility.",GREEN,17)
footer(s,4)

# 5
s=new_slide("4. Training the MultiVerS verifier","The first experiment trained the verifier on HealthVer.")
steps=[("Claim",1.0), ("Abstract",3.25), ("Longformer",5.5), ("Label + evidence heads",8.0), ("SUPPORT / CONTRADICT / NEI",10.45)]
for i,(txt,x) in enumerate(steps):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.25),Inches(1.75),Inches(1.05)); sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=TEAL
    textbox(s,x+.08,2.53,1.59,.45,txt,15,NAVY,True,PP_ALIGN.CENTER)
    if i<len(steps)-1:
        textbox(s,x+1.82,2.49,.45,.4,"→",25,TEAL,True,PP_ALIGN.CENTER)
bullets(s,["Original MultiVerS: trained for 20 epochs.","Fine-tuned MultiVerS: continued training from the best original checkpoint.","The final preferred verifier is the fine-tuned checkpoint.","No released checkpoint was used as the final result; the models were trained locally."],.9,4.15,11.6,1.8,18)
footer(s,5)

# 6
s=new_slide("5. Standalone verifier results","Evaluation on the HealthVer test set: 903 claim-document pairs.")
add_table(s,3,6,[["Model","Accuracy","Macro-F1","Abstract F1","Sentence selection F1","Sentence label F1"],["Original MultiVerS","74.86%","74.37%","74.90%","84.58%","73.52%"],["Fine-tuned MultiVerS","75.30%","74.89%","75.43%","85.41%","74.82%"]],.45,1.8,12.4,1.8,[2.5,1.7,1.7,2.0,3.0,2.3],12)
textbox(s,.85,4.2,11.7,1.0,"Conclusion: fine-tuning produced a small but consistent improvement. This checkpoint is frozen for the final MedGemma experiment.",22,GREEN,True,PP_ALIGN.CENTER)
footer(s,6)

# 7
s=new_slide("6. Model comparison experiment","A label-aware evidence selector was tested, but it was not better.")
add_table(s,4,4,[["Model","Accuracy","Macro-F1","Sentence selection F1"],["Original MultiVerS","74.86%","74.37%","84.58%"],["Fine-tuned MultiVerS","75.30%","74.89%","85.41%"],["Label-aware modified model","74.75%","74.37%","82.91%"]],1.2,1.7,10.9,2.2,[4.3,2.1,2.1,2.4],16)
textbox(s,1.25,4.45,10.8,1.1,"Decision: keep the fine-tuned original architecture as the best verifier. The next work improves retrieval and evaluation, not the model architecture.",21,ORANGE,True,PP_ALIGN.CENTER)
footer(s,7)

# 8
s=new_slide("7. Complete pipeline","The final system connects answer generation, claims, retrieval, and verification.")
flow=[("Medical question",.75,BLUE),("MedGemma answer",3.05,BLUE),("Atomic claims",5.35,TEAL),("Semantic retrieval",7.65,TEAL),("Frozen MultiVerS",9.95,GREEN)]
for i,(txt,x,col) in enumerate(flow):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.1),Inches(2.0),Inches(1.0)); sh.fill.solid(); sh.fill.fore_color.rgb=col; sh.line.fill.background(); textbox(s,x+.08,2.37,1.84,.42,txt,15,WHITE,True,PP_ALIGN.CENTER)
    if i<4: textbox(s,x+2.03,2.38,.3,.3,"→",22,TEAL,True,PP_ALIGN.CENTER)
textbox(s,2.0,4.25,9.4,.9,"Evidence-grounded answer\nwith SUPPORT / CONTRADICT / NEI and selected sentences",22,NAVY,True,PP_ALIGN.CENTER)
footer(s,8)

# 9
s=new_slide("8. Real MedGemma benchmark","This is the new final evaluation, using actual MedGemma-generated answers.")
add_table(s,4,2,[["Item","Number"],["Medical questions","60"],["Generated MedGemma answers","60"],["Extracted claims","187"]],1.0,1.7,5.2,2.5,[3.8,1.4],18)
card(s,6.8,1.7,5.3,3.3,"Benchmark coverage","Diseases, symptoms, treatments, medications, prevention, lifestyle, and diagnostics.\n\nEach claim was sent to biomedical semantic retrieval, then to the frozen fine-tuned MultiVerS model.",TEAL,17)
textbox(s,1.1,5.35,11,.55,"Retrieval model: S-PubMedBERT-MS-MARCO",18,NAVY,True,PP_ALIGN.CENTER)
footer(s,9)

# 10
s=new_slide("9. Current final-evaluation output","The system has produced predictions, but human labels are still required for scientific metrics.")
add_table(s,4,2,[["Automatic prediction","Count"],["SUPPORT","112"],["CONTRADICT","17"],["NEI","58"]],1.0,1.7,5.2,2.5,[3.8,1.4],18)
card(s,6.8,1.7,5.3,3.5,"Why accuracy is pending","The generated claims do not have gold human labels yet.\n\nA human reviewer must fill:\n• human label\n• human evidence\n\nThen Accuracy, Macro-F1, and evidence F1 can be calculated.",ORANGE,17)
footer(s,10)

# 11
s=new_slide("10. Manual evaluation and error analysis","The annotation file is ready for review.")
bullets(s,["File: results/medgemma_eval/manual_review.csv","Review whether each claim is supported by the retrieved evidence.","Add the correct human label: SUPPORT, CONTRADICT, or NEI.","Add the correct evidence document and sentence numbers.","Classify errors as: MedGemma hallucination, claim extraction error, retrieval error, or MultiVerS verification error."],.9,1.65,11.6,4.8,19)
footer(s,11)

# 12
s=new_slide("11. Current conclusion and next step","The project has progressed from dataset preparation to a complete evaluated pipeline.")
card(s,.9,1.7,5.5,3.8,"What is complete","✓ HealthVer converted and used for training\n✓ MultiVerS trained and fine-tuned\n✓ Best checkpoint selected\n✓ MedGemma integrated\n✓ Biomedical semantic retrieval integrated\n✓ 60-question real benchmark generated",GREEN,17)
card(s,6.9,1.7,5.5,3.8,"Next step","Manually annotate the 187 claims.\n\nThen calculate final metrics and compare:\n\nMedGemma only\nvs.\nMedGemma + fine-tuned MultiVerS",BLUE,17)
textbox(s,1.0,6.05,11.2,.45,"Main research question: Does evidence verification make MedGemma answers more reliable?",20,NAVY,True,PP_ALIGN.CENTER)
footer(s,12)

prs.save(OUT)
print(OUT)
